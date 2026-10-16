import io
import re
from typing import List, Dict, Any, Tuple
from docx import Document
from contract_app.infrastructure.storage.file_storage import file_storage
from contract_app.modules.revisions.domain.revision_plan import RevisionPlan

CLAUSE_REGEX = re.compile(
    r'^(?:'
    r'(第[一二三四五六七八九十百0-9]+[条章节款项])'
    r'|([一二三四五六七八九十百]+[、.．])'
    r'|([（\(][一二三四五六七八九十百0-9]+[）\)])'
    r'|(\d+(?:\.\d+)*[、.．]?)'
    r')\s*(.*)'
)

class DocxService:
    def extract_structure(self, object_key: str) -> List[Dict[str, Any]]:
        raw_bytes = file_storage.read_file(object_key)
        doc = Document(io.BytesIO(raw_bytes))
        blocks: List[Dict[str, Any]] = []

        para_idx = 0
        seen_texts = set()

        for p in doc.paragraphs:
            raw_text = p.text
            text = raw_text.strip()
            
            # Keep blank lines as spacing blocks (up to consecutive 1)
            if not text:
                if blocks and not blocks[-1].get("is_blank"):
                    para_idx += 1
                    blocks.append({
                        "block_id": f"blk_{para_idx}",
                        "clause_id": "",
                        "text": "",
                        "length": 0,
                        "align": "left",
                        "bold": False,
                        "indent": False,
                        "is_title": False,
                        "is_blank": True
                    })
                continue

            if text in seen_texts:
                continue
            seen_texts.add(text)
            para_idx += 1

            # Extract alignment: 1=CENTER, 2=RIGHT, 3=JUSTIFY, 0/None=LEFT
            align_val = p.alignment
            if align_val == 1 or (align_val is None and len(text) <= 25 and ("协议" in text or "合同" in text) and para_idx <= 2):
                align = "center"
            elif align_val == 2 or (align_val is None and text.startswith("协议编号")):
                align = "right"
            elif align_val == 3:
                align = "justify"
            else:
                align = "left"

            # Extract bold
            is_bold = any(bool(r.bold) for r in p.runs if r.text.strip())

            # Extract indent from docx: first_line_indent > 0
            if p.paragraph_format.first_line_indent is not None:
                has_indent = bool(p.paragraph_format.first_line_indent > 0)
            else:
                has_indent = raw_text.startswith(('  ', '\t', '　　', '    '))
                if not has_indent and len(text) > 35 and align not in ("center", "right") and not text.startswith(("甲方", "乙方", "法定代表人", "地址")):
                    has_indent = True

            # Detect document main title
            is_title = (align == "center" and len(text) <= 30 and para_idx <= 2)

            m = CLAUSE_REGEX.match(text)
            is_section_heading = False
            if m:
                matched_num = next(g for g in m.groups()[:4] if g is not None)
                heading_suffix = m.group(5).strip()
                if heading_suffix and len(heading_suffix) <= 20:
                    clause_id = f"{matched_num} {heading_suffix}"
                else:
                    clause_id = matched_num
                if any(c in matched_num for c in ['、', '章', '条']) and len(text) <= 40:
                    is_section_heading = True
            else:
                clause_id = f"段落 {para_idx}"

            if is_section_heading:
                is_bold = True
                if align not in ("center", "right"):
                    align = "left"

            blocks.append({
                "block_id": f"blk_{para_idx}",
                "clause_id": clause_id,
                "text": text,
                "length": len(text),
                "align": align,
                "bold": is_bold,
                "indent": has_indent,
                "is_title": is_title,
                "is_blank": False
            })

        # Also extract non-empty cells from tables if any
        if hasattr(doc, 'tables') and doc.tables:
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text and len(cell_text) > 10 and cell_text not in seen_texts:
                            seen_texts.add(cell_text)
                            para_idx += 1
                            blocks.append({
                                "block_id": f"blk_{para_idx}",
                                "clause_id": f"附表条款 {para_idx}",
                                "text": cell_text,
                                "length": len(cell_text),
                                "align": "left",
                                "bold": False,
                                "indent": False,
                                "is_title": False,
                                "is_blank": False
                            })

        return blocks

    def create_sample_docx(self, title: str, paragraphs: List[str]) -> Tuple[str, str]:
        doc = Document()
        doc.add_heading(title, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
        buf = io.BytesIO()
        doc.save(buf)
        raw = buf.getvalue()
        return file_storage.save_file(raw, f"{title}.docx")

    def apply_revision_plan(
        self,
        base_object_key: str,
        plan: RevisionPlan,
        clean: bool = False
    ) -> Tuple[str, str]:
        raw_bytes = file_storage.read_file(base_object_key)
        doc = Document(io.BytesIO(raw_bytes))

        # Build replacement mapping from operations
        replacements = {}
        for op in plan.operations:
            if clean and op.status == "rejected":
                continue
            if op.expected_text:
                replacements[op.expected_text.strip()] = op.replacement_text.strip()

        for p in doc.paragraphs:
            text = p.text
            for exp, rep in replacements.items():
                if exp in text:
                    p.text = text.replace(exp, rep)
                    text = p.text

        buf = io.BytesIO()
        doc.save(buf)
        raw = buf.getvalue()
        suffix = "_clean.docx" if clean else "_proposed.docx"
        return file_storage.save_file(raw, f"rev_{plan.id}{suffix}")

docx_service = DocxService()
